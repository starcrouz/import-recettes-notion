import os
import re
import urllib.parse
from bs4 import BeautifulSoup
from docx import Document
from docx.shared import Inches

# --- CONFIGURATION ---
INPUT_FILE = "mes_recettes.html"
OUTPUT_DIR = "Recettes_Word"

if not os.path.exists(OUTPUT_DIR):
    os.makedirs(OUTPUT_DIR)

def clean_filename(s):
    """Nettoie le titre pour le nom de fichier"""
    s = re.sub(r'[\\/*?:"<>|]', "", s)
    return s.strip().replace('\n', ' ')[:100]

print(f"lecture de {INPUT_FILE}...")

try:
    with open(INPUT_FILE, "r", encoding="utf-8") as f:
        soup = BeautifulSoup(f, "html.parser")
except FileNotFoundError:
    print(f"ERREUR : Le fichier '{INPUT_FILE}' est introuvable.")
    exit()

titres = soup.find_all("h2")
print(f" > {len(titres)} recettes détectées.")

for i, titre_balise in enumerate(titres):
    vrai_titre = titre_balise.get_text().strip()
    if not vrai_titre: continue

    nom_fichier = clean_filename(vrai_titre)
    doc = Document()
    
    # Analyse du contenu
    curr = titre_balise.next_sibling
    first_image_path = None
    elements_to_process = [] 
    
    while curr and curr.name != "h2":
        elements_to_process.append(curr)
        if first_image_path is None and curr.name in ['p', 'div', 'span']:
            imgs = curr.find_all('img')
            if imgs:
                src = imgs[0].get('src')
                if src:
                    src_decoded = urllib.parse.unquote(src)
                    if os.path.exists(src_decoded):
                        first_image_path = src_decoded
        curr = curr.next_sibling
        
    # --- CONSTRUCTION DU WORD ---
    
    # A. L'IMAGE en tout premier
    if first_image_path:
        try:
            doc.add_picture(first_image_path, width=Inches(5))
        except:
            pass

    # B. LE CORPS (Sans remettre le titre)
    image_already_placed = False 
    
    for elem in elements_to_process:
        if elem.name in ['p', 'ul', 'ol', 'div']:
            text = elem.get_text().strip()
            if text:
                doc.add_paragraph(text)
            
            images_in_elem = elem.find_all('img')
            for img in images_in_elem:
                src = img.get('src')
                if src:
                    src_decoded = urllib.parse.unquote(src)
                    if os.path.exists(src_decoded):
                        if src_decoded == first_image_path and not image_already_placed:
                            image_already_placed = True
                            continue 
                        try:
                            doc.add_picture(src_decoded, width=Inches(4))
                        except:
                            pass

    save_path = os.path.join(OUTPUT_DIR, nom_fichier + ".docx")
    try:
        doc.save(save_path)
        print(f"✅ Généré : {nom_fichier}")
    except Exception as e:
        print(f"❌ Erreur : {e}")

print("Terminé.")